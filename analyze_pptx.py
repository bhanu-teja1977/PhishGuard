import os
from pptx import Presentation

pptx_path = r"C:\Users\Bhanu Teja\.gemini\antigravity\scratch\PHISHGUARD\docs\PhishGuard.pptx"
pptx_text = []
if os.path.exists(pptx_path):
    prs = Presentation(pptx_path)
    for i, slide in enumerate(prs.slides):
        pptx_text.append(f"--- Slide {i+1} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                pptx_text.append(shape.text)

with open("pptx_full_text.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(pptx_text))
print("Extracted PPTX")
